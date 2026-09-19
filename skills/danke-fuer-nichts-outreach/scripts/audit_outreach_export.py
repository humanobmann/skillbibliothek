#!/usr/bin/env python3
"""Audit a CSV or XLSX outreach export for campaign safety risks."""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable


ALIASES = {
    "timestamp": ("timestamp", "sent_at", "sent time", "send time", "date", "datum", "zeit", "zeitpunkt"),
    "email": ("email", "recipient", "to", "empfänger", "empfaenger", "e-mail"),
    "subject": ("subject", "betreff"),
    "direction": ("direction", "richtung", "message_direction", "nachrichtenrichtung"),
    "diagnostic": ("diagnostic", "bounce", "bounce_reason", "smtp", "diagnose"),
    "attachments": ("attachments", "attachment", "attachment_names", "anhang", "anlagen", "anhänge", "anhaenge"),
    "content_type": ("content_type", "mime", "mime_type", "inhaltstyp"),
}


def norm(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def norm_email(value: Any) -> str:
    match = re.search(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", norm(value), re.I)
    return match.group(0).lower() if match else norm(value).lower()


def subject_signature(value: Any) -> str:
    text = norm(value).lower()
    text = re.sub(r"^.{1,40}:\s+(?=sortimentsanfrage|presse|rezension|gespräch|gespraech)", "{präfix}: ", text)
    text = re.sub(r"\b\d+\b", "#", text)
    return text


def parse_time(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        parsed = value
        return parsed.astimezone(timezone.utc).replace(tzinfo=None) if parsed.tzinfo else parsed
    text = norm(value)
    if not text:
        return None
    text = text.replace("Z", "+00:00")
    for candidate in (text, text.replace(".", "-")):
        try:
            parsed = datetime.fromisoformat(candidate)
            return parsed.astimezone(timezone.utc).replace(tzinfo=None) if parsed.tzinfo else parsed
        except ValueError:
            pass
    for fmt in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def classify_bounce(value: Any) -> str:
    text = norm(value).lower()
    if not text:
        return "NONE"
    if re.search(r"\b5\.7\.\d+\b|spam|policy|reputation|as\(4810\)", text):
        return "SPAM_BLOCK"
    if re.search(r"\b5\.[01]\.\d+\b|user unknown|mailbox unavailable|does not exist|invalid recipient", text):
        return "HARD_INVALID"
    if re.search(r"\b4\.\d\.\d+\b|temporar|mailbox full|try again", text):
        return "SOFT_BOUNCE"
    return "REVIEW"


def resolve_headers(headers: Iterable[str]) -> dict[str, str | None]:
    lookup = {norm(h).lower(): h for h in headers if h is not None}
    resolved: dict[str, str | None] = {}
    for key, aliases in ALIASES.items():
        resolved[key] = next((lookup[a.lower()] for a in aliases if a.lower() in lookup), None)
    return resolved


def load_rows(path: Path, sheet: str | None) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise SystemExit("XLSX benötigt openpyxl. CSV ist ohne Zusatzpaket möglich.") from exc
        book = load_workbook(path, read_only=True, data_only=True)
        ws = book[sheet] if sheet else book[book.sheetnames[0]]
        values = ws.iter_rows(values_only=True)
        headers = [norm(v) for v in next(values)]
        return [dict(zip(headers, row)) for row in values]
    raise SystemExit("Unterstützt werden CSV, XLSX und XLSM.")


def max_in_window(times: list[datetime], minutes: int) -> int:
    times = sorted(times)
    left = 0
    best = 0
    span = timedelta(minutes=minutes)
    for right, current in enumerate(times):
        while current - times[left] > span:
            left += 1
        best = max(best, right - left + 1)
    return best


def audit(rows: list[dict[str, Any]], max_day: int, max_window: int) -> dict[str, Any]:
    if not rows:
        return {"status": "STOP", "errors": ["Keine Datenzeilen gefunden."]}
    headers = resolve_headers(rows[0].keys())
    missing = [name for name in ("timestamp", "email", "subject") if not headers[name]]
    if missing:
        return {"status": "STOP", "errors": ["Fehlende Pflichtspalten: " + ", ".join(missing)]}

    outbound_values = {"ausgang", "outbound", "outgoing", "sent", "gesendet"}
    outreach_rows = (
        [row for row in rows if norm(row.get(headers["direction"] or "")).lower() in outbound_values]
        if headers["direction"]
        else rows
    )
    if headers["direction"] and not outreach_rows:
        return {"status": "STOP", "errors": ["Richtungsspalte erkannt, aber keine Ausgangszeilen gefunden."]}

    emails: list[str] = []
    subjects: list[str] = []
    times: list[datetime] = []
    day_counts: Counter[str] = Counter()
    bounce_counts: Counter[str] = Counter()
    tnef_rows = 0
    unparsed_times = 0
    attachment_values: list[str] = []

    for row in rows:
        diagnostic = row.get(headers["diagnostic"] or "") if headers["diagnostic"] else ""
        bounce_counts[classify_bounce(diagnostic)] += 1

    for row in outreach_rows:
        email = norm_email(row.get(headers["email"] or ""))
        if email:
            emails.append(email)
        subjects.append(subject_signature(row.get(headers["subject"] or "")))
        timestamp = parse_time(row.get(headers["timestamp"] or ""))
        if timestamp:
            times.append(timestamp)
            day_counts[timestamp.date().isoformat()] += 1
        else:
            unparsed_times += 1
        payload = " ".join(
            norm(row.get(headers[key] or "")) for key in ("attachments", "content_type") if headers[key]
        ).lower()
        if headers["attachments"]:
            attachment_values.append(norm(row.get(headers["attachments"] or "")).lower())
        if "winmail.dat" in payload or "application/ms-tnef" in payload:
            tnef_rows += 1

    email_counts = Counter(emails)
    subject_counts = Counter(s for s in subjects if s)
    max_day_count = max(day_counts.values(), default=0)
    max_10m = max_in_window(times, 10)
    repeated_subject = subject_counts.most_common(1)[0] if subject_counts else ("", 0)
    risks: list[str] = []
    warnings: list[str] = []

    if max_day_count > max_day:
        risks.append(f"Tagesmaximum {max_day_count} liegt über {max_day}.")
    if max_10m > max_window:
        risks.append(f"Zehn-Minuten-Maximum {max_10m} liegt über {max_window}.")
    if bounce_counts["SPAM_BLOCK"]:
        risks.append(f"{bounce_counts['SPAM_BLOCK']} Spam- oder Policyblocks gefunden.")
    if tnef_rows:
        risks.append(f"{tnef_rows} Zeilen mit TNEF oder winmail.dat gefunden.")
    if repeated_subject[1] >= 5 and repeated_subject[1] / max(len(outreach_rows), 1) >= 0.3:
        risks.append(f"Serielles Betreffmuster in {repeated_subject[1]} von {len(outreach_rows)} Ausgangszeilen.")
    duplicate_events = sum(count - 1 for count in email_counts.values() if count > 1)
    if duplicate_events:
        warnings.append(f"{duplicate_events} zusätzliche Versandereignisse an wiederholte Empfänger.")
    if unparsed_times:
        warnings.append(f"{unparsed_times} Zeitwerte konnten nicht gelesen werden.")
    flag_values = {"", "true", "false", "1", "0", "yes", "no", "ja", "nein"}
    if attachment_values and set(attachment_values).issubset(flag_values) and not headers["content_type"]:
        warnings.append("Die Anlagenspalte enthält nur Ja/Nein-Werte. TNEF oder winmail.dat kann damit nicht ausgeschlossen werden.")
    if bounce_counts["REVIEW"]:
        warnings.append(f"{bounce_counts['REVIEW']} Bounce-Diagnosen benötigen manuelle Prüfung.")

    return {
        "status": "STOP" if risks else ("CHECK" if warnings else "READY"),
        "source_rows": len(rows),
        "rows": len(outreach_rows),
        "unique_recipients": len(set(emails)),
        "duplicate_send_events": duplicate_events,
        "max_per_day": max_day_count,
        "max_per_10_minutes": max_10m,
        "most_common_subject_signature": {"value": repeated_subject[0], "count": repeated_subject[1]},
        "bounce_classes": dict(sorted(bounce_counts.items())),
        "tnef_rows": tnef_rows,
        "risks": risks,
        "warnings": warnings,
        "resolved_headers": headers,
    }


def to_markdown(result: dict[str, Any]) -> str:
    lines = [
        f"# Outreach-Audit: {result.get('status', 'STOP')}",
        "",
        f"* Zeilen: {result.get('rows', 0)}",
        f"* eindeutige Empfänger: {result.get('unique_recipients', 0)}",
        f"* Tagesmaximum: {result.get('max_per_day', 0)}",
        f"* Maximum in zehn Minuten: {result.get('max_per_10_minutes', 0)}",
        f"* TNEF oder winmail.dat: {result.get('tnef_rows', 0)}",
        "",
    ]
    for heading, key in (("Sperrgründe", "risks"), ("Prüfhinweise", "warnings"), ("Fehler", "errors")):
        values = result.get(key, [])
        if values:
            lines.extend([f"## {heading}", "", *[f"* {value}" for value in values], ""])
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", type=Path)
    parser.add_argument("--sheet")
    parser.add_argument("--max-per-day", type=int, default=8)
    parser.add_argument("--max-per-10-minutes", type=int, default=8)
    parser.add_argument("--format", choices=("json", "markdown"), default="markdown")
    parser.add_argument("--fail-on-stop", action="store_true")
    args = parser.parse_args()
    result = audit(load_rows(args.path, args.sheet), args.max_per_day, args.max_per_10_minutes)
    print(json.dumps(result, ensure_ascii=False, indent=2) if args.format == "json" else to_markdown(result))
    return 2 if args.fail_on_stop and result.get("status") == "STOP" else 0


if __name__ == "__main__":
    sys.exit(main())
